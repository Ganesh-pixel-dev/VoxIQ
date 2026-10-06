"""Read a resume (PDF, DOCX or TXT) into numbered lines and pull out project names and skills."""
import io
import re
from dataclasses import dataclass, field

BULLETS = "•●▪◦-*·–"
SECTIONS = {
    "projects": ("project", "projects", "academic projects", "personal projects", "key projects"),
    "experience": ("experience", "work experience", "internship", "internships", "employment", "professional experience"),
    "skills": ("skills", "technical skills", "key skills", "tech stack", "technologies"),
    "education": ("education", "academics", "qualifications"),
    "other": ("summary", "objective", "profile", "achievements", "certifications", "awards", "interests", "languages", "about"),
}


class ResumeError(ValueError):
    pass


@dataclass
class Resume:
    lines: list                                    # [{"n": 1, "text": "...", "section": "projects", "heading": False}]
    projects: list = field(default_factory=list)   # [{"name": ..., "line": n}]
    skills: list = field(default_factory=list)


def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if name.endswith(".txt"):
            return data.decode("utf-8", errors="replace")
    except Exception as e:
        raise ResumeError(
            f"Could not read this file as a resume ({type(e).__name__}). Export it again or upload it as .txt."
        ) from e
    raise ResumeError("Resume must be a .pdf, .docx or .txt file.")


def _heading(line: str):
    if len(line) > 40:
        return None
    key = re.sub(r"[^a-z ]", "", line.lower()).strip()
    for section, names in SECTIONS.items():
        if key in names:
            return section
    return None


def _is_bullet(text: str) -> bool:
    return text[:1] in BULLETS


def _project_name(text: str) -> str:
    name = re.split(r"\s[|–—-]\s|:\s|\(|\s{2,}", text, maxsplit=1)[0]
    return name.strip(" " + BULLETS)


def parse(text: str) -> Resume:
    lines, section = [], "header"
    for raw in text.splitlines():
        t = raw.strip()
        if not t:
            continue
        h = _heading(t)
        if h:
            section = h
        lines.append({"n": len(lines) + 1, "text": t, "section": section, "heading": bool(h)})
    if len([l for l in lines if not l["heading"]]) < 3:
        raise ResumeError("I could not find enough text in this file. Scanned images are not supported (no OCR).")

    projects, skills = [], []
    for l in lines:
        if l["heading"]:
            continue
        if l["section"] == "projects" and not _is_bullet(l["text"]) and len(l["text"].split()) <= 14:
            name = _project_name(l["text"])
            if 2 <= len(name) <= 60:
                projects.append({"name": name, "line": l["n"]})
        if l["section"] == "skills":
            body = l["text"].split(":", 1)[-1]
            skills += [s.strip(" " + BULLETS) for s in re.split(r"[,;|/]", body) if 1 < len(s.strip()) <= 30]
    seen, uniq = set(), []
    for s in skills:
        if s.lower() not in seen:
            seen.add(s.lower())
            uniq.append(s)
    return Resume(lines=lines, projects=projects, skills=uniq)
